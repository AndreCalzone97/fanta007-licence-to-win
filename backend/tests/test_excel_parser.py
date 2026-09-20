from collections import Counter
from pathlib import Path

import pytest

from app.data.excel_parser import parse_players


@pytest.fixture
def official_raw_players():
    source = Path(__file__).resolve().parents[2] / "data/source/Quotazioni_Fantacalcio_Stagione_2026_27_2026-09-03.xlsx"
    if not source.is_file():
        pytest.skip("Official quotation export unavailable; synthetic parser tests still run")
    return parse_players(source)


@pytest.mark.official_exports
def test_parser_reads_complete_active_list(official_raw_players):
    raw_players = official_raw_players
    assert len(raw_players) == 533
    assert Counter(player.role_classic for player in raw_players) == {
        "P": 65,
        "D": 188,
        "C": 193,
        "A": 87,
    }
    assert len({player.id for player in raw_players}) == 533


@pytest.mark.official_exports
def test_parser_reads_locatelli_from_real_dataset(official_raw_players):
    raw_players = official_raw_players
    locatelli = next(player for player in raw_players if player.id == 827)
    assert locatelli.source_name == "Locatelli"
    assert locatelli.team == "Juventus"
    assert locatelli.role_classic == "C"
    assert locatelli.current_quotation == 9
    assert locatelli.fvm == 33
    assert locatelli.fvm_mantra == 38


@pytest.mark.official_exports
def test_all_quotation_deltas_are_consistent(official_raw_players):
    raw_players = official_raw_players
    assert all(
        player.quotation_delta
        == player.current_quotation - player.initial_quotation
        for player in raw_players
    )
    assert all(
        player.quotation_delta_mantra
        == player.current_quotation_mantra - player.initial_quotation_mantra
        for player in raw_players
    )


def test_parser_reads_synthetic_workbook(quotation_source, raw_players):
    parsed = parse_players(quotation_source)
    assert parsed == raw_players
    assert len(parsed) == 9
    assert Counter(row.role_classic for row in parsed) == {"P": 1, "D": 2, "C": 2, "A": 4}
    assert len({row.id for row in parsed}) == 9


@pytest.mark.parametrize("case", ["duplicate", "classic_delta", "mantra_delta", "role", "integer", "text", "header", "sheet", "empty"])
def test_parser_rejects_invalid_workbooks(quotation_source, case):
    from openpyxl import load_workbook
    from app.data.excel_parser import DatasetValidationError

    workbook = load_workbook(quotation_source)
    sheet = workbook["Tutti"]
    if case == "duplicate":
        sheet["A4"] = sheet["A3"].value
    elif case == "classic_delta":
        sheet["H3"] = 99
    elif case == "mantra_delta":
        sheet["K3"] = 99
    elif case == "role":
        sheet["B3"] = "X"
    elif case == "integer":
        sheet["F3"] = "invalid"
    elif case == "text":
        sheet["D3"] = " "
    elif case == "header":
        sheet["A2"] = "Wrong ID"
    elif case == "sheet":
        sheet.title = "Wrong sheet"
    else:
        sheet.delete_rows(3, sheet.max_row)
    workbook.save(quotation_source)
    workbook.close()

    with pytest.raises(DatasetValidationError):
        parse_players(quotation_source)
